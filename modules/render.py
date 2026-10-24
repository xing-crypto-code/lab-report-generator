"""
docx → HTML 渲染器 · render.py
=============================
把生成的 .docx 渲染成可在网页预览的 HTML（含表格结构、正文、虚线占位框、脚注占位），
让用户无需下载即可看到报告实际效果。
"""
import html

_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# 占位框尺寸（与 generator 保持同一口径）
_PH_H, _PH_W = 6.5, 12.0


def _tbl_is_dashed(tbl_el) -> bool:
    """该表格是否是虚线占位框（四边 dashed）"""
    tblPr = tbl_el.find(_NS + 'tblPr')
    if tblPr is None:
        return False
    b = tblPr.find(_NS + 'tblBorders')
    if b is None:
        return False
    vals = []
    for edge in ('top', 'left', 'bottom', 'right'):
        el = b.find(_NS + edge)
        if el is not None:
            vals.append(el.get(_NS + 'val'))
    return bool(vals) and all(v == 'dashed' for v in vals)


def _tbl_dims(tbl_el):
    """读取表格宽/行高（cm），读不到就给默认占位尺寸"""
    w_cm, h_cm = _PH_W, _PH_H
    try:
        tblPr = tbl_el.find(_NS + 'tblPr')
        if tblPr is not None:
            tw = tblPr.find(_NS + 'tblW')
            if tw is not None and tw.get(_NS + 'w'):
                w_cm = int(tw.get(_NS + 'w')) / 567.0
    except Exception:
        pass
    try:
        tr = tbl_el.find(_NS + 'tr')
        if tr is not None:
            trPr = tr.find(_NS + 'trPr')
            if trPr is not None:
                h = trPr.find(_NS + 'trHeight')
                if h is not None and h.get(_NS + 'val'):
                    h_cm = int(h.get(_NS + 'val')) / 567.0
    except Exception:
        pass
    return w_cm, h_cm


def _all_text(el) -> str:
    return "".join(t.text or "" for t in el.findall('.//' + _NS + 't')).strip()


def _dashed_box_html(tbl_el) -> str:
    """把虚线占位表格渲染成网页上的虚线占位框"""
    label = _all_text(tbl_el) or '【此处插入图片】'
    w_cm, h_cm = _tbl_dims(tbl_el)
    return (
        f"<div class='r-ph' style='width:{w_cm:.1f}cm;min-height:{h_cm:.1f}cm'>"
        f"<span>{html.escape(label)}</span></div>"
    )


def _cell_html(cell) -> str:
    """把一个单元格渲染为 HTML（含内部嵌套表格 / 虚线占位框）"""
    inner = []
    tc = cell._tc
    for child in tc:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            texts = []
            for r in child.findall('.//' + _NS + 'r'):
                t = r.find(_NS + 't')
                if t is not None and t.text:
                    texts.append(t.text)
                else:
                    draw = r.find('.//' + _NS + 'drawing')
                    if draw is not None:
                        texts.append('🖼️[图片]')
            txt = "".join(texts).strip()
            if txt:
                inner.append(f"<div class='r-para'>{html.escape(txt)}</div>")
        elif tag == 'tbl':
            if _tbl_is_dashed(child):
                inner.append(_dashed_box_html(child))
            else:
                from docx.table import Table
                nst = Table(child, cell)
                inner.append(_table_html(nst, nested=True))
    return "".join(inner) if inner else "&nbsp;"


def _table_html(table, nested=False) -> str:
    cls = "r-table" + (" r-nested" if nested else "")
    ncols = len(table.columns) if table.columns else 1
    rows = []
    # 处理合并单元格：用 row.cells 但去重
    for row in table.rows:
        cells = row.cells
        seen = []
        prev = None
        for c in cells:
            key = id(c._tc)
            if prev is not None and key == prev[0]:
                continue
            seen.append((key, c))
            prev = (key, c)
        tds = "".join(f"<td>{_cell_html(c)}</td>" for _, c in seen)
        rows.append(f"<tr>{tds}</tr>")
    thead = "".join(f"<col style='width:auto'/>" for _ in range(ncols))
    return f"<table class='{cls}'>{thead}{''.join(rows)}</table>"


def render_docx(docx_path: str) -> str:
    """把 docx 渲染为 HTML 字符串"""
    from docx import Document
    doc = Document(docx_path)
    out = ["<div class='r-page'>"]

    # 逐 body 元素渲染（保持段落与表格顺序）
    body = doc.element.body
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    tbl_idx = 0
    for child in body:
        tag = child.tag
        if tag == ns + 'p':
            texts = []
            for r in child.findall('.//' + ns + 'r'):
                t = r.find(ns + 't')
                if t is not None and t.text:
                    texts.append(t.text)
            txt = "".join(texts).strip()
            if txt:
                out.append(f"<div class='r-para'>{html.escape(txt)}</div>")
        elif tag == ns + 'tbl':
            if _tbl_is_dashed(child):
                out.append(_dashed_box_html(child))
                tbl_idx += 1
                continue
            from docx.table import Table
            tb = Table(child, doc)
            if tbl_idx < len(doc.tables):
                tb = doc.tables[tbl_idx]
            out.append(_table_html(tb))
            tbl_idx += 1

    out.append("</div>")
    return "".join(out)