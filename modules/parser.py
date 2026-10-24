"""
文档解析模块 · parser.py
职责：识别用户上传的实验指导手册（PDF / 图片），
      提取其中的文字、表格、实验步骤、数据等信息。

OCR 引擎：EasyOCR（支持 80+ 语言，中文识别优秀，Python 3.14 兼容）
"""
import fitz  # PyMuPDF
from PIL import Image
from docx import Document
import io
import os
from pathlib import Path
from typing import Optional

from config import OCR_LANG, OCR_USE_GPU
from modules.preprocessor import preprocess_for_ocr, detect_document_region

def _detect_cuda() -> bool:
    """
    自动检测当前 torch 是否支持 CUDA（GPU 加速）。
    只有装了 CUDA 版 torch 才返回 True，否则强制 CPU，
    避免 EasyOCR 在无 CUDA 时报错。
    """
    try:
        import torch
        return bool(torch.cuda.is_available())
    except Exception:
        return False


# GPU 可用性自动检测（覆盖 .env 的 OCR_USE_GPU）
_CUDA_OK = _detect_cuda()
if OCR_USE_GPU and not _CUDA_OK:
    print("[parser] 检测到 OCR_USE_GPU=true，但当前 torch 为 CPU 版本，"
          f"已自动降级为 CPU 模式。可尝试安装 CUDA 版 torch 使用 GPU 加速。")

# EasyOCR 延迟加载（首次调用时初始化，避免启动卡顿）
_reader = None
_reader_error = None


def ocr_models_ready() -> bool:
    """检查 EasyOCR 模型文件是否已下载到本地。

    EasyOCR 首次使用时会去 GitHub 下载模型；国内网络下常被阻断，
    表现是**无限等待**（用户看到界面一直转圈）。提前检查可快速失败。
    """
    model_dir = os.path.expanduser('~/.EasyOCR/model')
    need = ['craft_mlt_25k.pth']
    if OCR_LANG == "ch":
        need.append('zh_sim_g2.pth')
    missing = [f for f in need if not os.path.exists(os.path.join(model_dir, f))]
    if missing:
        return False
    return True


def _get_reader():
    """单例模式获取 EasyOCR Reader 实例。

    模型未就绪时**直接抛错**而不是挂起等下载——否则整个生成流程会卡死。
    """
    global _reader, _reader_error
    if _reader is not None:
        return _reader
    if _reader_error is not None:
        raise RuntimeError(_reader_error)

    if not ocr_models_ready():
        _reader_error = (
            "OCR 模型尚未下载完成（首次使用需下载识别模型）。\n"
            "本机无法直连 GitHub，请让程序自动走镜像下载，或改用 PDF/Word 格式的指导书。"
        )
        raise RuntimeError(_reader_error)

    try:
        import easyocr
        # EasyOCR 语言代码: ch_sim=简体中文, en=英文
        lang_list = ["ch_sim", "en"] if OCR_LANG == "ch" else ["en"]
        use_gpu = OCR_USE_GPU and _CUDA_OK
        _reader = easyocr.Reader(
            lang_list=lang_list,
            gpu=use_gpu,
            verbose=False,
            download_enabled=False,   # 禁止运行时下载，避免卡死
        )
    except Exception as e:
        _reader_error = f"OCR 初始化失败：{e}"
        raise RuntimeError(_reader_error)
    return _reader


def parse_pdf_markdown(file_path: str) -> str:
    """用 pymupdf4llm 把 PDF 转成结构化 Markdown（标题/列表/表格都保留）。

    这是比裸文本 OCR 精确得多的抽取方式：它保留版面层级，标题会变成 `#`，
    列表会保留 `1.` `-` 编号，表格转成 Markdown 表格。让 AI 能分清
    「实验目的」「实验仪器」这类小节的边界，从而原样照抄而不是重新编。
    失败时返回空串，调用方回退到原来的 parse_pdf。
    """
    try:
        import pymupdf4llm
        md = pymupdf4llm.to_markdown(file_path)
        return md or ""
    except Exception as e:
        print(f"[parser] pymupdf4llm 提取失败，回退裸文本: {e}")
        return ""


def parse_pdf(file_path: str) -> dict:
    """
    解析 PDF 文件，返回：
    {
        "text": "整段文字内容",
        "markdown": "结构化 Markdown（供精确抽取小节用）",
        "tables": [["表头", "列2"], ["数据1", "数据2"]],
        "page_count": 页数
    }
    """
    # ── 先尝试结构化 Markdown 抽取（电子版 PDF 效果最好）──
    markdown = parse_pdf_markdown(file_path)

    doc = fitz.open(file_path)
    full_text = []
    all_tables = []
    page_count = len(doc)

    for page_num in range(page_count):
        page = doc[page_num]

        # 1. 提取文字（电子版 PDF）
        text = page.get_text("text")
        if text.strip():
            full_text.append(f"--- 第 {page_num+1} 页 ---\n{text}")

        # 2. 电子版没文字（扫描件）才走 OCR，避免重复且慢
        if not text.strip():
            images = page.get_images(full=True)
            for img_index, img in enumerate(images):
                xref = img[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                pil_image = Image.open(io.BytesIO(image_bytes))

                ocr_result = ocr_image(pil_image)
                if ocr_result["text"]:
                    full_text.append(
                        f"--- 第 {page_num+1} 页 · 图片 {img_index+1} ---\n"
                        f"{ocr_result['text']}"
                    )
                all_tables.extend(ocr_result["tables"])

    doc.close()
    return {
        "text": "\n".join(full_text),
        "markdown": markdown,
        "tables": all_tables,
        "page_count": page_count
    }


# OCR 保护参数：CPU 版 torch 下 EasyOCR 很慢（单张高清图可达数分钟），
# 这里限制输入尺寸并设硬超时，避免整个生成流程被一张图拖死。
_OCR_MAX_SIDE = 1600      # 长边超过此值先等比缩小（OCR 精度基本无损，速度提升数倍）
_OCR_TIMEOUT = 90         # 单张图 OCR 硬超时（秒）


def _shrink_for_ocr(image: Image.Image) -> Image.Image:
    """把长边过大的图等比缩小，避免 OCR 在超大图上耗时爆炸。"""
    try:
        w, h = image.size
        longest = max(w, h)
        if longest > _OCR_MAX_SIDE:
            scale = _OCR_MAX_SIDE / float(longest)
            new_size = (max(1, int(w * scale)), max(1, int(h * scale)))
            return image.resize(new_size, Image.LANCZOS)
    except Exception:
        pass
    return image


def _ocr_with_timeout(reader, img_bgr, timeout=_OCR_TIMEOUT):
    """在独立线程里跑 OCR，超时则放弃该图（返回空结果而非卡死整个流程）。"""
    import threading
    box = {}

    def _work():
        try:
            box['result'] = reader.readtext(img_bgr, detail=1)
        except Exception as e:
            box['error'] = e

    th = threading.Thread(target=_work, daemon=True)
    th.start()
    th.join(timeout)
    if th.is_alive():
        print(f"[parser] OCR 超时（>{timeout}s），跳过该图")
        return []
    if 'error' in box:
        print(f"[parser] OCR 出错：{box['error']}")
        return []
    return box.get('result', [])


def ocr_image(image: Image.Image) -> dict:
    """
    对单张图片做 OCR 识别（含图像预处理增强）
    返回 { "text": "识别文字", "tables": [表格列表] }
    """
    import numpy as np

    # ── 先缩小超大图，防止 CPU 版 OCR 耗时爆炸 ──
    image = _shrink_for_ocr(image)

    # ── 预处理 ──
    processed = preprocess_for_ocr(image)

    # PIL → OpenCV 格式（RGB → BGR）
    img_array = np.array(processed.convert("RGB"))
    img_bgr = img_array[:, :, ::-1].copy()

    reader = _get_reader()
    results = _ocr_with_timeout(reader, img_bgr)

    texts = []
    tables = []

    for item in results:
        bbox, text, confidence = item
        if confidence > 0.3:
            texts.append(_fix_ocr_typos(text))

    return {
        "text": "\n".join(texts),
        "tables": tables
    }


# 中文 OCR 常见的固定误识（字形相近），做一次确定性纠正。
# 只纠正**无歧义**的高频错字，不动可能合法的内容。
_OCR_TYPO_MAP = [
    ('熟恶', '熟悉'), ('熟悉', '熟悉'),
    ('夹验', '实验'), ('实脸', '实验'), ('实骏', '实验'),
    ('与非门', '与非门'),
    ('74L500', '74LS00'), ('74LSO0', '74LS00'), ('74LS0O', '74LS00'),
    ('74LS0', '74LS0'), ('74L520', '74LS20'), ('74L532', '74LS32'),
    ('74L502', '74LS02'), ('74L586', '74LS86'),
    ('数模电', '数模电'),
]


def _fix_ocr_typos(text: str) -> str:
    """纠正 OCR 高频误识的固定字串。"""
    s = text or ""
    for bad, good in _OCR_TYPO_MAP:
        if bad != good and bad in s:
            s = s.replace(bad, good)
    # 逻辑表达式里的常见符号误识：A:B -> A·B
    s = s.replace('(A:B)', '(A·B)')
    return s


def parse_image(file_path: str) -> dict:
    """解析单张图片文件（照片/截图）"""
    image = Image.open(file_path)
    try:
        return ocr_image(image)
    except RuntimeError as e:
        raise ValueError(
            f"{e}\n"
            "👉 建议：把指导书另存为 PDF 或 Word 上传，这两类格式不需要 OCR、秒出结果。"
        )


def parse_docx(file_path: str) -> dict:
    """
    解析 .docx/.doc Word 文档的实验指导手册
    （有些老师发的是 Word 版指导书）
    """
    # .doc（老格式）尝试用 python-docx，失败则报友好错误
    ext = Path(file_path).suffix.lower()
    if ext == ".doc":
        try:
            doc = Document(file_path)
        except Exception:
            raise ValueError(
                "暂不支持 .doc 旧格式 Word 文档。"
                "请用 WPS/Word 打开后「另存为」→ 选择 .docx 格式再上传。"
            )
    else:
        doc = Document(file_path)

    full_text = []
    all_tables = []

    # 提取段落文本
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            full_text.append(text)

    # 提取表格
    for table in doc.tables:
        rows_data = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_data.append(cells)
        if rows_data:
            all_tables.append(rows_data)

    return {
        "text": "\n".join(full_text),
        "tables": all_tables,
        "page_count": None  # docx 没有明确页数概念
    }


def parse_markdown(file_path: str) -> dict:
    """
    解析 .md / .txt / .markdown 纯文本&Markdown 实验指导手册。
    Markdown 里的表格（| a | b |）会被提取为结构化表格。
    """
    try:
        text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
    except Exception:
        text = Path(file_path).read_text(encoding="gbk", errors="ignore")

    lines = text.splitlines()
    clean_lines = []
    all_tables = []
    cur_table = []
    in_table = False

    def flush_table():
        nonlocal cur_table
        # 过滤纯分隔行 |---|
        rows = [r for r in cur_table if not _is_md_sep(r)]
        if rows:
            all_tables.append(
                [[c.strip() for c in r.strip(" |").split("|")] for r in rows]
            )
        cur_table = []

    for ln in lines:
        stripped = ln.strip()
        if _is_md_sep(stripped):
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                cur_table = []
            cur_table.append(stripped)
            clean_lines.append(stripped)
        else:
            if in_table:
                flush_table()
                in_table = False
            if stripped and not stripped.lstrip().startswith((
                    "#", "```", ">", "-", "*", "|", "<!--")):
                clean_lines.append(stripped)
    if in_table:
        flush_table()

    return {
        "text": "\n".join(clean_lines),
        "tables": all_tables,
        "page_count": None
    }


def _is_md_sep(s: str) -> bool:
    """判断是否 Markdown 表格分隔行（| --- | :---: | --- |）"""
    if not (s.strip().startswith("|") and s.strip().endswith("|")):
        return False
    cells = [c.strip() for c in s.strip().strip("|").split("|")]
    if not cells:
        return False
    # 所有单元格都只能由 -、:、空白组成，且至少含一个 -
    return all(c != "" and all(ch in "-:" + " " for ch in c) and "-" in c for c in cells)


def parse_document(file_path: str) -> dict:
    """统一入口：自动判断格式，返回统一结构"""
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in {".docx", ".doc"}:
        return parse_docx(file_path)
    elif ext in {".md", ".markdown", ".txt"}:
        return parse_markdown(file_path)
    elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"}:
        return parse_image(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


# ---------- 便捷测试 ----------
if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        result = parse_document(sys.argv[1])
        print(f"页数: {result.get('page_count', 'N/A')}")
        print(f"文字长度: {len(result['text'])} 字符")
        print(f"表格数: {len(result['tables'])}")
        print("--- 前 500 字预览 ---")
        print(result["text"][:500])