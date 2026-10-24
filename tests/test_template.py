"""
模板回归测试 · tests/test_template.py
=====================================
守护内置模板的排版修复不被回退。每新增一条 TEMPLATE_FIXES，就同步加一条断言。

运行：python tests/test_template.py
"""
import os
import sys
import warnings

warnings.filterwarnings('ignore')
for _k in ('HTTP_PROXY', 'HTTPS_PROXY', 'http_proxy', 'https_proxy',
           'ALL_PROXY', 'all_proxy', 'NO_PROXY', 'no_proxy'):
    os.environ.pop(_k, None)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, 'libs'))

from docx import Document                       # noqa: E402
from docx.oxml.ns import qn                     # noqa: E402

TEMPLATE = os.path.join(ROOT, 'templates', '内置实验报告模板.docx')

_failures = []


def check(name, cond, detail=''):
    if cond:
        print(f'  ✓ {name}')
    else:
        print(f'  ✗ {name}  {detail}')
        _failures.append(name)


def main():
    if not os.path.exists(TEMPLATE):
        print(f'✗ 模板不存在: {TEMPLATE}')
        return 1
    doc = Document(TEMPLATE)

    print('\n== fix_01 表格宽度自洽（tblW == ΣgridCols == ΣtcW）==')
    for ti, t in enumerate(doc.tables):
        grid = t._tbl.find(qn('w:tblGrid'))
        check(f'table{ti} 有 tblGrid', grid is not None)
        if grid is None:
            continue
        cols = [int(g.get(qn('w:w')) or 0) for g in grid.findall(qn('w:gridCol'))]
        total = sum(cols)
        tblW = t._tbl.tblPr.find(qn('w:tblW'))
        check(f'table{ti} 有 tblW', tblW is not None)
        check(f'table{ti} tblW type=dxa',
              tblW is not None and tblW.get(qn('w:type')) == 'dxa')
        check(f'table{ti} tblW({tblW.get(qn("w:w")) if tblW is not None else None}) == ΣgridCols({total})',
              tblW is not None and int(tblW.get(qn('w:w')) or 0) == total)
        for ri, row in enumerate(t.rows):
            seen = set()
            s = 0
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                tcPr = c._tc.find(qn('w:tcPr'))
                if tcPr is None:
                    continue
                w = tcPr.find(qn('w:tcW'))
                if w is not None and w.get(qn('w:w')):
                    s += int(w.get(qn('w:w')))
            check(f'table{ti} r{ri} ΣtcW({s}) == ΣgridCols({total})', s == total)

    print('\n== fix_02 无表格级居中 ==')
    for ti, t in enumerate(doc.tables):
        check(f'table{ti} 无 tblPr/jc',
              t._tbl.tblPr.find(qn('w:jc')) is None)

    print('\n== fix_03 顶层表格等宽（上下表边框对齐）==')
    widths = []
    for t in doc.tables:
        grid = t._tbl.find(qn('w:tblGrid'))
        if grid is None:
            continue
        widths.append(sum(int(g.get(qn('w:w')) or 0)
                          for g in grid.findall(qn('w:gridCol'))))
    check(f'所有顶层表格等宽: {sorted(set(widths))}', len(set(widths)) == 1)

    print('\n== fix_04 单元格边距规范（左右108/上下0）==')
    for ti, t in enumerate(doc.tables):
        mar = t._tbl.tblPr.find(qn('w:tblCellMar'))
        check(f'table{ti} 有 tblCellMar', mar is not None)
        if mar is None:
            continue
        for edge, want in (('left', 108), ('right', 108), ('top', 0), ('bottom', 0)):
            el = mar.find(qn('w:' + edge))
            got = int(el.get(qn('w:w'))) if el is not None and el.get(qn('w:w')) else None
            check(f'table{ti} {edge}={got} (期望{want})', got == want)

    print('\n== fix_05 Normal 样式宋体小四 ==')
    st = doc.styles['Normal']
    check(f'Normal 字体={st.font.name}', st.font.name == '宋体')
    check(f'Normal 字号={st.font.size.pt if st.font.size else None}pt（期望12.0）',
          st.font.size is not None and abs(st.font.size.pt - 12.0) < 0.01)

    print('\n== fix_06 表格横向定位统一（上下表左右边框重合）==')
    sec = doc.sections[0]
    body_left = int(sec.left_margin / 635)
    edges = []
    for ti, t in enumerate(doc.tables):
        tblPr = t._tbl.tblPr
        grid = t._tbl.find(qn('w:tblGrid'))
        total = sum(int(g.get(qn('w:w')) or 0) for g in grid.findall(qn('w:gridCol')))
        tblpPr = tblPr.find(qn('w:tblpPr'))
        anchor = tblpPr.get(qn('w:horzAnchor')) if tblpPr is not None else None
        px = tblpPr.get(qn('w:tblpX')) if tblpPr is not None else None
        xspec = tblpPr.get(qn('w:tblpXSpec')) if tblpPr is not None else None
        check(f'table{ti} 无 tblpXSpec（不居中/不右对齐）', xspec is None)
        check(f'table{ti} 锚点=margin', anchor == 'margin')
        check(f'table{ti} tblpX=0（贴版心左边缘）', px in (None, '0'))
        left = body_left + int(px or 0)
        edges.append((left, left + total))
    check(f'所有表左边界一致: {sorted({e[0] for e in edges})}',
          len({e[0] for e in edges}) == 1)
    check(f'所有表右边界一致: {sorted({e[1] for e in edges})}',
          len({e[1] for e in edges}) == 1)
    check(f'表格级无 jc',
          all(t._tbl.tblPr.find(qn('w:jc')) is None for t in doc.tables))

    print('\n== fix_07 填充时不产生模板外的多余章节 ==')
    import tempfile
    from modules.generator import fill_template
    _secs = {
        '实验目的': '测试目的', '实验仪器': '测试仪器',
        '实验原理': '测试原理',
        '实验过程': '一、实验步骤\n1. 测试。',
        '实验小结': '测试小结',
        # 下面这些模板里没有，必须被丢弃
        '实验设计': '不应出现', '思考题': '不应出现',
    }
    _out = None
    try:
        _out = fill_template(TEMPLATE, {'sections': _secs, 'tables_fill': []},
                             '_regress_extra')
        _d = Document(_out)
        body_text = '\n'.join((p.text or '') for p in _d.paragraphs)
        check('未在正文新增"实验设计"标题', '实验设计' not in body_text)
        check('未在正文新增"思考题"标题', '思考题' not in body_text)
        check('未新增任何多余顶层段落',
              not [p for p in _d.paragraphs
                   if (p.text or '').strip() and (p.text or '').strip() != '成都大学实验报告'])
    finally:
        if _out and os.path.exists(_out):
            os.remove(_out)

    print('\n== fix_07 表格布局 fixed（autofit 会让边框随内容漂移）==')
    for ti, t in enumerate(doc.tables):
        lay = t._tbl.tblPr.find(qn('w:tblLayout'))
        check(f'table{ti} tblLayout=fixed',
              lay is not None and lay.get(qn('w:type')) == 'fixed')

    print('\n== fix_08 任意模板填充后左右边框都对齐版心 ==')
    check('版心宽度=9780', int((sec.page_width - sec.left_margin - sec.right_margin) / 635) == 9780)

    print('\n== fix_08 pPr/rPr 子元素顺序合法 + 无异常缩进 ==')
    PPR = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
           'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
           'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct',
           'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
           'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents',
           'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
           'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
           'sectPr', 'pPrChange']
    RPR = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps',
           'strike', 'dstrike', 'outline', 'shadow', 'emboss', 'imprint',
           'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing',
           'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect',
           'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang',
           'eastAsianLayout', 'specVanish', 'oMath']

    def _sorted_ok(el, order):
        tags = [c.tag.split('}')[-1] for c in el]
        idxs = [order.index(t) for t in tags if t in order]
        return idxs == sorted(idxs)

    bad_ppr = [p for p in doc.element.body.iter(qn('w:pPr')) if not _sorted_ok(p, PPR)]
    check(f'pPr 顺序合法（错误 {len(bad_ppr)} 个）', not bad_ppr)
    bad_rpr = [r for r in doc.element.body.iter(qn('w:rPr')) if not _sorted_ok(r, RPR)]
    check(f'rPr 顺序合法（错误 {len(bad_rpr)} 个）', not bad_rpr)

    abn = 0
    for p in doc.element.body.iter(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        if pPr is None:
            continue
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            continue
        fc = ind.get(qn('w:firstLineChars'))
        if fc is not None:
            try:
                if abs(int(fc)) > 10:
                    abn += 1
            except Exception:
                pass
    check(f'无异常首行缩进（firstLineChars>10 的段落 {abn} 个）', abn == 0)

    print('\n== fix_09 填充结果无多余尾部空段（Delete 不再吃上一行）==')
    _secs2 = {
        '实验目的': '测试目的', '实验仪器': '测试仪器',
        '实验原理': '一、推导\n两段测试文字。\n第二段测试文字。',
        '实验过程': '一、实验步骤\n第一步。\n第二步。',
        '实验小结': '测试小结',
    }
    _out2 = None
    try:
        _out2 = fill_template(TEMPLATE, {'sections': _secs2, 'tables_fill': []},
                              '_regress_blank')
        _d2 = Document(_out2)
        _t2 = _d2.tables[1]
        blanks = 0
        for _row in _t2.rows:
            _seen = set()
            for _c in _row.cells:
                if id(_c._tc) in _seen:
                    continue
                _seen.add(id(_c._tc))
                _ps = [_x for _x in _c._tc if _x.tag == qn('w:p')]
                if _ps and not any((t.text or '').strip()
                                   for t in _ps[-1].iter(qn('w:t'))):
                    # 末尾空段只允许存在于完全空的行（如教师评阅栏）
                    if (_c.text or '').strip():
                        blanks += 1
        check(f'有内容的单元格无尾部空段（{blanks} 个异常）', blanks == 0)
    finally:
        if _out2 and os.path.exists(_out2):
            os.remove(_out2)

    print('\n== fix_10 编号格式统一（1. 2. 3.），保证左边缘对齐 ==')
    import re as _re
    from modules.generator import _iter_render_items
    _mixed = ('1.第一条说明\n2、第二条说明\n（3）第三条说明\n'
              '④第四条说明\n5. 第五条说明')
    _items = _iter_render_items(_mixed)
    _num = [i for i in _items if i["kind"] == "list"]
    _seps = set()
    _pat = _re.compile(r'^(?:\d+\s*([\.、．)）])|[（(]\d+[）)]|[①②③④⑤⑥⑦⑧⑨⑩]|[a-zA-Z][\.、)])')
    for _it in _num:
        _m = _pat.match(_it["text"])
        if _m:
            _seps.add(_m.group(1) or 'other')
    check(f'混用编号已归一（分隔符 {sorted(_seps)}）', _seps <= {'.'})
    check(f'编号条目数=5（实际{len(_num)}）', len(_num) == 5)
    check('编号从 1 连续递增',
          [t["text"].split(".")[0] for t in _num] == ['1', '2', '3', '4', '5'])

    print('\n== fix_11 结构化原文提取（pymupdf4llm）==')
    try:
        import pymupdf4llm  # noqa: F401
        check('pymupdf4llm 已安装', True)
    except Exception as _e:
        check(f'pymupdf4llm 已安装（{_e}）', False)
    try:
        from modules.parser import parse_pdf_markdown
        check('parse_pdf_markdown 可导入', callable(parse_pdf_markdown))
    except Exception as _e:
        check(f'parse_pdf_markdown 可导入（{_e}）', False)

    print('\n== 版式参数（A4 + 成都大学边距）==')
    sec = doc.sections[0]
    check(f'页面 {sec.page_width.cm:.1f}x{sec.page_height.cm:.1f}cm',
          abs(sec.page_width.cm - 21.0) < 0.1 and abs(sec.page_height.cm - 29.7) < 0.1)
    check(f'左边距 {sec.left_margin.cm:.2f}cm', abs(sec.left_margin.cm - 1.75) < 0.05)
    check(f'右边距 {sec.right_margin.cm:.2f}cm', abs(sec.right_margin.cm - 2.00) < 0.05)

    print('\n== 主表结构（7行：目的/仪器/原理/过程/小结/评阅）==')
    main_t = doc.tables[1] if len(doc.tables) > 1 else None
    check('存在主表', main_t is not None)
    if main_t is not None:
        check(f'主表 7 行（实际{len(main_t.rows)}）', len(main_t.rows) == 7)
        labels = [(main_t.rows[i].cells[0].text or '').strip() for i in range(len(main_t.rows))]
        for want in ('实验目的', '实验仪器', '实验原理', '实验过程', '实验小结'):
            check(f'含「{want}」栏', any(want in lb for lb in labels))

    print('\n' + '=' * 50)
    if _failures:
        print(f'✗ {len(_failures)} 项未通过:')
        for f in _failures:
            print('   -', f)
        return 1
    print('✓ 全部通过')
    return 0


if __name__ == '__main__':
    sys.exit(main())
